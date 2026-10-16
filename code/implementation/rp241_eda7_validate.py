"""Independent EDA7 acceptance checks against immutable original outputs."""
from pathlib import Path
import json,hashlib,zipfile,csv,re
import pandas as pd,numpy as np
from pandas.testing import assert_frame_equal
from docx import Document
R=Path('H:/SCI2/MRQI');D='2026-10-16';E=R/f'3_results/rp241_eda6_manuscript_supplements_{D}';O=R/f'3_results/rp241_eda7_methods_alignment_{D}';Q=R/f'0_admin/validation/rp241_eda7_{D}';M=R/f'5_manuscript/rp241_eda7_methods_results_alignment_{D}'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def load(p):return pd.read_csv(p,sep='\t')
parent=R/f'6_release/phase_packages/MRQI_RP241_EDA6_results_complete_supplements_DUC_unsigned_delta_{D}.zip';assert sha(parent)=='65ba340738dfa90b4656788fe4dad1fe735b4cf856a01c6e4a34e0b78b622f31'
with zipfile.ZipFile(parent) as z:
 assert z.testzip() is None and len(z.namelist())==77
 rows=list(csv.DictReader(z.read(next(n for n in z.namelist() if n.startswith('0_admin/manifest/'))).decode().splitlines(),delimiter='\t'))
 for r in rows:assert sha(R/r['path'])==r['sha256'] and (R/r['path']).stat().st_size==int(r['bytes'])
provenance=load(O/f'Input_Provenance_Receipt_{D}.tsv')
for r in provenance.itertuples():assert (R/r.path).stat().st_size==r.bytes and sha(R/r.path)==r.sha256
idx=load(E/f'Supplementary_Table_Index_{D}.tsv');assert len(idx)==24 and idx.rows.sum()==121573
for r in idx.itertuples():assert sha(E/r.file)==r.sha256 and load(E/r.file).shape==(r.rows,r.columns)
a=load(E/'S04a_All81_Complete_Evidence.tsv');b=load(E/'S04b_Original43_Complete_Evidence.tsv');n=load(O/f'S14a_All81_Aligned_Complete_Evidence_{D}.tsv');p=load(O/f'S14b_Original43_Aligned_Complete_Evidence_{D}.tsv')
assert n.shape==(81,163) and p.shape==(43,163) and n.HGNC_ID.is_unique
assert_frame_equal(n[a.columns],a,check_dtype=False,rtol=3e-14,atol=2e-14)
assert_frame_equal(p[b.columns].reset_index(drop=True),b,check_dtype=False,rtol=3e-14,atol=2e-14)
assert set(p.HGNC_ID)==set(a.loc[a.primary43,'HGNC_ID']) and len(n)==81
protein=load(E/'S09_Public_Protein_All81_All8_Source_Models.tsv');ad=protein[protein.original_analysis_file.eq('AD_PRIMARY_All81_NoPrune.tsv')].set_index('HGNC_ID').reindex(n.HGNC_ID)
for c in ['status','n','cases','controls','OLS_beta','OLS_SE','OLS_CI_low','OLS_CI_high','OLS_P','OLS_Q_primary43','OLS_Q_all81_exploratory','HC3_P','HC3_Q_primary43','expected_beta_sign']:
 x=n['AD_public_protein_'+c].reset_index(drop=True);y=ad[c].reset_index(drop=True)
 pd.testing.assert_series_equal(x,y,check_names=False,check_dtype=False,rtol=3e-14,atol=2e-14)
# Null, opposite and missing observations must remain explicit, at both scopes.
sign=p.direction.str.startswith('R1').map({True:1,False:-1})
assert p.BA17_beta.notna().sum()==41 and (p.BA17_beta*sign>0).sum()==29 and (p.BA17_beta*sign<0).sum()==12 and p.BA17_q43.lt(.05).sum()==0
assert set(p.loc[p.BA17_beta.isna(),'gene'])=={'OR7A5','PRPF4'}
assert p.GBM_both_all_donors_agree.sum()==17 and p.GBM_both_all_LODO_same_peak.sum()==32
assert a.GBM132_beta.isna().sum()==21 and a.AD159_BY508.lt(.05).sum()==20
pp=protein[protein.primary43];aa=pp[pp.original_analysis_file.eq('AD_PRIMARY_All81_NoPrune.tsv')].set_index('HGNC_ID');gg=pp[pp.original_analysis_file.eq('GBM_PRIMARY_All81_NoPrune.tsv')].set_index('HGNC_ID').reindex(aa.index)
ag=aa.OLS_beta*aa.expected_beta_sign>0;bg=gg.OLS_beta*gg.expected_beta_sign>0
assert aa.OLS_P.notna().sum()==24 and gg.OLS_P.notna().sum()==36
assert set(aa.loc[ag&bg&aa.OLS_Q_primary43.lt(.05)&gg.OLS_Q_primary43.lt(.05),'HGNC_symbol'])=={'PSME3','CKAP4','NRIP2'}
assert (ag&bg&aa.HC3_Q_primary43.lt(.05)&gg.HC3_Q_primary43.lt(.05)).sum()==0
assert set(gg.loc[~bg&gg.OLS_Q_primary43.lt(.05),'HGNC_symbol'])=={'ARC','SUN2'}
cross=load(O/f'S12_Manuscript_Section_Figure_Table_Alignment_Index_{D}.tsv');assert len(cross)==55 and not cross.duplicated(['item_type','item_id']).any()
assert cross.item_type.value_counts().to_dict()=={'Supplement':24,'Methods':14,'Legend':9,'Results':8}
assert set(cross.loc[cross.item_type.eq('Supplement'),'item_id'])==set(idx.table_id)
assert cross.original_statistics_changed.eq(False).all() and cross.methods.notna().all()
reg=load(O/f'S13_Source_Model_Family_Alignment_Registry_{D}.tsv');assert len(reg)==124 and not reg.duplicated(['source','model']).any()
assert len(reg[reg.source.eq('GSE174367')])==36 and len(reg[reg.source.eq('GSE159699')])==24
ba=reg[reg.source.eq('GSE203206')];assert ba.N.astype(int).tolist()==[46,46,37,35,27,27,46]
assert reg[reg.source.eq('GSE237183')].N.astype(int).tolist()==[2,4,3]
assert set(reg[reg.source.str.startswith('CPTAC')|reg.source.str.startswith('PXD')].primary_family.astype(int))==set(int(x.OLS_P.notna().sum()) for _,x in protein[protein.primary43].groupby('original_analysis_file'))
assert reg.model.eq('S_CNV_complete').sum()==1 and reg.model.eq('S_CNV_concordant').sum()==1
payload=json.loads((M/f'Methods_Aligned_Results_Legends_Source_{D}.json').read_text(encoding='utf-8'))
assert len(payload['methods'])==14 and len(payload['results'])==8 and len(payload['legends'])==9
old=json.loads((R/f'5_manuscript/rp241_eda6_results_{D}/Results_Text_And_Legends_Editable_Source_{D}.json').read_text(encoding='utf-8'))
for (oldtitle,oldbody),(title,body) in zip(old['results'],payload['results']):
 assert title==oldtitle
 # All previously reported scientific numbers must still occur, in sequence.
 nums=re.findall(r'(?<![A-Za-z])\d+(?:\.\d+)?',oldbody)
 newnums=re.findall(r'(?<![A-Za-z])\d+(?:\.\d+)?',body);it=iter(newnums);assert all(any(x==target for x in it) for target in nums)
docx=next(M.glob('*.docx'));doc=Document(docx);text='\n'.join(x.text for x in doc.paragraphs)
for token in ['cooksCutoff=FALSE','23/28','12/15','32,136','24,448','EX的S2仅11AD/4','未进行跨脑区随机效应元分析','S_CNV_concordant','已有DUC仅预填','任意基因依赖','无一同时保留两来源HC3支持']:assert token in text,token
assert len(doc.tables)==1 and len(doc.tables[0].rows)==29
render=json.loads((Q/f'docx_render/WPS_Render_Receipt_{D}.json').read_text(encoding='utf-8-sig'));assert render['source_sha256']==sha(docx) and render['visible'] is False and render['libreoffice_used'] is False
pages=json.loads((Q/'docx_render/Page_Text_And_Counts.json').read_text(encoding='utf-8'));assert len(pages)==7 and all(r['chars']>900 for r in pages)
fm=load(R/f'4_figures/rp241_eda6_{D}/Figure_Manifest_{D}.tsv');assert len(fm)==9
for r in fm.itertuples():assert sha(R/r.relative_path)==r.sha256
result=dict(date=D,status='PASS_EDA7_METHODS_RESULTS_LEGENDS_TABLES_AND_IDENTITY_ALIGNMENT',Methods_sections=14,Results_sections=8,legends=9,original_supplement_tables=24,original_supplement_records=121573,crosswalk_items=55,source_model_context_rows=124,complete81_fields=n.shape[1],original43_fields=p.shape[1],BA17_main_FDR_support=0,public_protein_joint_HC3_primary_support=0,negative_conflict_missing_rows_preserved=True,all_76_EDA6_assets_unchanged=True,figure_files_unchanged=9,docx_pages=7,new_models=0,new_P_or_q=0,controlled_data_read=False,author_approval_claimed=False)
(Q/f'Independent_Alignment_Validation_{D}.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');print(json.dumps(result,ensure_ascii=False,indent=2))
